from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


COLOR_PRINCIPAL = "234A72"
COLOR_SECUNDARIO = "EAF0F6"
COLOR_BORDE = "C9D3DE"


def _escribir_fila(hoja, valores, fila):
    for columna, valor in enumerate(valores, start=1):
        celda = hoja.cell(row=fila, column=columna, value=valor)
        if isinstance(valor, str):
            celda.data_type = "s"
    return fila + 1


def _estilizar_encabezado(hoja, fila, cantidad_columnas):
    relleno = PatternFill("solid", fgColor=COLOR_PRINCIPAL)
    fuente = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    borde = Border(bottom=Side(style="thin", color=COLOR_BORDE))
    for celda in hoja[fila][:cantidad_columnas]:
        celda.fill = relleno
        celda.font = fuente
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde
    hoja.row_dimensions[fila].height = 24


def _ajustar_columnas(hoja):
    for columna in hoja.columns:
        indice = columna[0].column
        ancho = max(
            (len(str(celda.value)) for celda in columna if celda.value is not None),
            default=10,
        )
        hoja.column_dimensions[get_column_letter(indice)].width = min(max(ancho + 3, 12), 38)


def exportar_reporte_pagos_excel(ruta, datos):
    """Genera un libro con el detalle de pagos y el resumen del período."""
    libro = Workbook()
    pagos = libro.active
    pagos.title = "Detalle de pagos"
    pagos.freeze_panes = "A2"

    encabezados = [
        "ID pago",
        "Fecha y hora",
        "Huésped",
        "Identificación",
        "Habitación",
        "ID reserva",
        "Monto",
        "Método de pago",
        "Registrado por",
    ]
    _escribir_fila(pagos, encabezados, 1)
    _estilizar_encabezado(pagos, 1, len(encabezados))

    for fila, pago in enumerate(datos["pagos"], start=2):
        valores = [
            pago["pago_id"],
            pago["creado_en"],
            pago["cliente"],
            pago["identificacion"],
            pago["habitacion_numero"],
            pago["reserva_id"],
            float(pago["monto"]),
            pago["metodo_pago"],
            pago["empleado_nombre"],
        ]
        _escribir_fila(pagos, valores, fila)
        pagos.cell(fila, 2).number_format = "dd/mm/yyyy hh:mm"
        pagos.cell(fila, 7).number_format = '"$"#,##0.00'
        if fila % 2 == 0:
            for celda in pagos[fila]:
                celda.fill = PatternFill("solid", fgColor=COLOR_SECUNDARIO)
    if pagos.max_row > 1:
        pagos.auto_filter.ref = pagos.dimensions
    _ajustar_columnas(pagos)

    resumen = libro.create_sheet("Resumen")
    resumen.merge_cells("A1:C1")
    resumen["A1"] = "REPORTE DE PAGOS"
    resumen["A1"].font = Font(size=16, bold=True, color=COLOR_PRINCIPAL)
    resumen["A2"] = "Período"
    resumen["B2"] = datos["fecha_inicio"]
    resumen["C2"] = datos["fecha_fin"]
    resumen["B2"].number_format = "dd/mm/yyyy"
    resumen["C2"].number_format = "dd/mm/yyyy"
    resumen["A4"] = "Cantidad de pagos"
    resumen["B4"] = datos["cantidad_pagos"]
    resumen["A5"] = "Total recibido"
    resumen["B5"] = float(datos["total"])
    resumen["B5"].number_format = '"$"#,##0.00'

    fila_encabezado = 7
    resumen.cell(fila_encabezado, 1, "Fecha")
    resumen.cell(fila_encabezado, 2, "Cantidad de pagos")
    resumen.cell(fila_encabezado, 3, "Total recibido")
    _estilizar_encabezado(resumen, fila_encabezado, 3)
    for fila, total_dia in enumerate(datos["por_dia"], start=fila_encabezado + 1):
        resumen.cell(fila, 1, total_dia["fecha"]).number_format = "dd/mm/yyyy"
        resumen.cell(fila, 2, total_dia["cantidad"])
        celda_total = resumen.cell(fila, 3, float(total_dia["total"]))
        celda_total.number_format = '"$"#,##0.00'

    fila_metodos = fila_encabezado + len(datos["por_dia"]) + 4
    resumen.merge_cells(
        start_row=fila_metodos,
        start_column=1,
        end_row=fila_metodos,
        end_column=3,
    )
    resumen.cell(fila_metodos, 1, "Ingresos por método de pago")
    resumen.cell(fila_metodos, 1).font = Font(
        size=13, bold=True, color=COLOR_PRINCIPAL
    )
    fila_metodos += 1
    for columna, encabezado in enumerate(
        ["Método de pago", "Cantidad de pagos", "Total recibido"], start=1
    ):
        resumen.cell(fila_metodos, columna, encabezado)
    _estilizar_encabezado(resumen, fila_metodos, 3)
    for fila, metodo in enumerate(datos["por_metodo"], start=fila_metodos + 1):
        resumen.cell(fila, 1, metodo["metodo_pago"])
        resumen.cell(fila, 2, metodo["cantidad"])
        celda_total = resumen.cell(fila, 3, float(metodo["total"]))
        celda_total.number_format = '"$"#,##0.00'
    _ajustar_columnas(resumen)

    libro.properties.title = "Reporte de pagos del hotel"
    libro.properties.subject = (
        f"Pagos del {datos['fecha_inicio']} al {datos['fecha_fin']}"
    )
    libro.save(Path(ruta))
